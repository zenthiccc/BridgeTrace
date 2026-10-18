"""Shaping sheet, report naming, and the two-sheet report workbook."""
from datetime import date, datetime, time
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

import config
import shaping
from convert_calendar import convert
from main import report_filename, run_pipeline

# Raw export order (End Date before Start Date) plus a column the shaping step drops.
RAW_HEADERS = [
    "Department", "Employee ID", "Employee Name", "Internal Ref", "Leave Code", "End Date",
    "Start Date", "Last Modified", "Start Time", "End Time", "Leave Days", "Leave Hours",
]


def raw_row(emp_id, code, start, changed=date(2026, 1, 20), start_time=time(0, 0)):
    return {
        "Department": "DEPT-A",
        "Employee ID": emp_id,
        "Employee Name": "Test User",
        "Internal Ref": "X",
        "Leave Code": code,
        "End Date": datetime.combine(start, time()),
        "Start Date": datetime.combine(start, time()),
        "Last Modified": datetime.combine(changed, time()),
        "Start Time": start_time,
        "End Time": time(0, 0),
        "Leave Days": 1,
        "Leave Hours": 8,
    }


def frame(*rows):
    return pd.DataFrame(list(rows), columns=RAW_HEADERS)


def write_raw(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(RAW_HEADERS)
    for row in rows:
        ws.append([row[h] for h in RAW_HEADERS])
    wb.save(path)


@pytest.fixture
def calendar_file(tmp_path):
    path = tmp_path / "calendar.xlsx"
    convert(config.CALENDAR_MD, path)
    return path


# --- Shaping step ----------------------------------------------------------

def test_transform_keeps_target_leave_in_column_order():
    out = shaping.transform(frame(
        raw_row("E1", "SL", date(2026, 1, 12)),
        raw_row("E1", "AL", date(2026, 1, 13)),  # filtered out
        raw_row("E2", "SLX", date(2026, 1, 14)),
    ))
    assert list(out.columns) == [*config.SHAPE_COLUMNS, *config.SHAPE_ADDED_COLUMNS]
    assert out["Leave Code"].tolist() == ["SL", "SLX"]


def test_transform_applies_types():
    row = shaping.transform(frame(
        raw_row("E1", "SL", date(2026, 8, 4), start_time=time(8, 30)),
    )).iloc[0]
    assert row["Start Date"] == pd.Timestamp(2026, 8, 4)
    assert row["Start Time"] == time(8, 30)
    assert all(row[c] is None for c in config.SHAPE_ADDED_COLUMNS)


def test_transform_lists_newest_change_first_keeping_ties_in_order():
    out = shaping.transform(frame(
        raw_row("E1", "SL", date(2026, 1, 5), changed=date(2026, 1, 6)),
        raw_row("E2", "SL", date(2026, 1, 7), changed=date(2026, 2, 1)),
        raw_row("E3", "SLH", date(2026, 1, 8), changed=date(2026, 1, 6)),
    ))
    assert out["Employee ID"].tolist() == ["E2", "E1", "E3"]


def test_transform_reports_missing_columns():
    raw = frame(raw_row("E1", "SL", date(2026, 1, 5))).drop(columns=["Leave Hours"])
    with pytest.raises(ValueError, match="Leave Hours"):
        shaping.transform(raw)


# --- Report naming ---------------------------------------------------------

@pytest.mark.parametrize("raw, expected", [
    ("leave_export_20260803_075106.xlsx", "Continuous Leave Report 2026-08-03.xlsx"),
    (r"C:\exports\leave_export_20261231.xlsx", "Continuous Leave Report 2026-12-31.xlsx"),
])
def test_report_filename_uses_export_date(raw, expected):
    assert report_filename(Path(raw)) == expected


def test_report_filename_without_date_uses_today():
    assert report_filename("leave data.xlsx") == f"Continuous Leave Report {date.today():%Y-%m-%d}.xlsx"


# --- Report workbook -------------------------------------------------------

def test_run_pipeline_writes_both_sheets(tmp_path, calendar_file):
    raw = tmp_path / "leave_export_20260113.xlsx"
    write_raw(raw, [
        raw_row("E1", "SL", date(2026, 1, 9)),   # Fri
        raw_row("E1", "SL", date(2026, 1, 12)),  # Mon -> weekend bridge
        raw_row("E2", "AL", date(2026, 1, 12)),  # not on the first sheet
    ])
    out = tmp_path / report_filename(raw)

    summary = run_pipeline(raw, calendar_file, out, log=lambda _msg: None)

    assert summary == (2, 1)
    wb = load_workbook(out)
    assert wb.sheetnames == [config.SHEET_FILTERED, config.SHEET_CONTINUOUS]

    sheet = wb[config.SHEET_FILTERED]
    assert [c.value for c in sheet[1]] == [*config.SHAPE_COLUMNS, *config.SHAPE_ADDED_COLUMNS]
    assert sheet.tables[config.FILTERED_TABLE].ref == "A1:M3"
    assert sheet["E2"].number_format == "m/d/yyyy"            # Start Date
    assert sheet["H2"].number_format.startswith("[$-409]h:mm:ss")  # Start Time

    assert list(wb[config.SHEET_CONTINUOUS].iter_rows(min_row=2, values_only=True)) == [(
        "E1", "Test User", datetime(2026, 1, 9), datetime(2026, 1, 12), 4,
        "Sick Leave -> Weekend -> Sick Leave",
    )]


def test_run_pipeline_without_target_leave_still_writes_a_valid_table(tmp_path, calendar_file):
    raw = tmp_path / "leave_export_20260113.xlsx"
    write_raw(raw, [raw_row("E2", "AL", date(2026, 1, 12))])
    out = tmp_path / "report.xlsx"

    assert run_pipeline(raw, calendar_file, out, log=lambda _msg: None) == (0, 0)
    sheet = load_workbook(out)[config.SHEET_FILTERED]
    assert sheet.tables[config.FILTERED_TABLE].ref == "A1:M2"
