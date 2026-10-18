"""Pipeline entry point.

Usage:
    python src/main.py [path-to-raw.xlsx] [path-to-output.xlsx]

If no raw file is given, the newest .xlsx in data/raw/ is used. If no output path is
given, the report is saved in data/output/ and dated from the raw file name
(any YYYYMMDD in it, e.g. leave_export_20260803.xlsx).
"""
from __future__ import annotations

import math
import re
import sys
from datetime import date, datetime
from pathlib import Path
from typing import NamedTuple

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import SortCondition, SortState
from openpyxl.worksheet.table import Table, TableStyleInfo

import config
import shaping
from calendar_utils import WorkCalendar
from load_data import load_records
from models import Sequence
from sequence_detector import detect_all

DATE_FORMAT = "m/d/yyyy"
TIME_FORMAT = r"[$-409]h:mm:ss\ AM/PM;@"
HIGHLIGHT = PatternFill("solid", fgColor="FFE699")

_EXPORT_DATE_RE = re.compile(r"(?<!\d)(\d{4})(\d{2})(\d{2})(?!\d)")


class ReportSummary(NamedTuple):
    target_rows: int  # rows on the filtered sheet
    sequences: int    # rows on the continuous-leave sheet


def _find_raw_file() -> Path:
    candidates = sorted(config.DATA_RAW.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    if not candidates:
        raise FileNotFoundError(
            f"No .xlsx found in {config.DATA_RAW}. Place the raw export there or "
            f"pass a path: python src/main.py <raw.xlsx>"
        )
    return candidates[0]


def export_date(raw_path) -> date | None:
    """First valid YYYYMMDD in a raw export's file name."""
    for year, month, day in _EXPORT_DATE_RE.findall(Path(raw_path).stem):
        try:
            return date(int(year), int(month), int(day))
        except ValueError:
            continue
    return None


def report_filename(raw_path) -> str:
    """Report file name for a raw export, dated today if its name has no date."""
    return config.REPORT_FILENAME.format(export_date(raw_path) or date.today())


def _chain_text(tokens: list[str]) -> str:
    """Render tokens as readable text, e.g. 'Sick Leave -> Weekend -> Sick Leave'."""
    parts = []
    for token in tokens:
        if token in config.LEAVE_TEXT:
            parts.append(config.LEAVE_TEXT[token])
        elif token in ("weekend", "holiday"):
            parts.append(token.capitalize())
        else:
            parts.append(token)
    return " -> ".join(parts)


def to_dataframe(sequences: list[Sequence]) -> pd.DataFrame:
    columns = [
        config.OUT_EMP_ID,
        config.OUT_EMP_NAME,
        config.OUT_START_DATE,
        config.OUT_END_DATE,
        config.OUT_CONSECUTIVE_DAYS,
        config.OUT_LEAVE_CHAIN,
    ]
    rows = [
        (s.emp_id, s.emp_name, s.start_date, s.end_date, s.consecutive_days, _chain_text(s.tokens))
        for s in sequences
    ]
    return pd.DataFrame(rows, columns=columns)


def _cell_value(value):
    """Plain Python value for openpyxl: missing -> None, pandas/numpy scalars unwrapped."""
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value.item() if hasattr(value, "item") else value


def _write_rows(ws, df: pd.DataFrame) -> None:
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append([_cell_value(v) for v in row])


def _write_filtered_sheet(ws, df: pd.DataFrame) -> None:
    """Shaped export as a formatted Excel table, sorted newest change first."""
    _write_rows(ws, df)
    headers = list(df.columns)
    last_row = max(ws.max_row, 2)  # an Excel table needs at least one data row
    formats = {"date": DATE_FORMAT, "time": TIME_FORMAT}

    for idx, name in enumerate(headers, start=1):
        kind = config.SHAPE_COLUMNS.get(name, "text")
        ws.cell(row=1, column=idx).fill = HIGHLIGHT
        if kind in formats or name == config.SHAPE_SORT_COLUMN:
            for (cell,) in ws.iter_rows(min_row=2, max_row=last_row, min_col=idx, max_col=idx):
                if kind in formats:
                    cell.number_format = formats[kind]
                if name == config.SHAPE_SORT_COLUMN:
                    cell.fill = HIGHLIGHT
        longest = 11 if kind in formats else max((len(str(v)) for v in df[name].dropna()), default=0)
        # +4 leaves room for the header's filter button.
        ws.column_dimensions[get_column_letter(idx)].width = min(max(len(name) + 4, longest + 2), 50)

    end_col = get_column_letter(len(headers))
    table = Table(
        displayName=config.FILTERED_TABLE,
        ref=f"A1:{end_col}{last_row}",
        tableStyleInfo=TableStyleInfo(showRowStripes=True),
    )
    if len(df):
        sort_col = get_column_letter(headers.index(config.SHAPE_SORT_COLUMN) + 1)
        table.sortState = SortState(
            ref=f"A2:{end_col}{last_row}",
            sortCondition=[SortCondition(ref=f"{sort_col}2:{sort_col}{last_row}", descending=True)],
        )
    ws.add_table(table)


def _write_sequences_sheet(ws, df: pd.DataFrame) -> None:
    """Sequences with dates as M/D/YYYY and wrapped, auto-sized columns and rows."""
    _write_rows(ws, df)

    max_width = 60
    wrap = Alignment(wrap_text=True, vertical="top")
    ncols, nrows = ws.max_column, ws.max_row
    header = {ws.cell(row=1, column=c).value: c for c in range(1, ncols + 1)}
    date_cols = {header[n] for n in (config.OUT_START_DATE, config.OUT_END_DATE) if n in header}

    def display_text(cell, col_idx) -> str:
        val = cell.value
        if val is None:
            return ""
        if col_idx in date_cols and isinstance(val, (date, datetime)):
            return f"{val.month}/{val.day}/{val.year}"
        return str(val)

    widths: dict[int, float] = {}
    for c in range(1, ncols + 1):
        longest = 0
        for r in range(1, nrows + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = wrap
            if r == 1:
                cell.font = Font(bold=True)
            elif c in date_cols and isinstance(cell.value, (date, datetime)):
                cell.number_format = DATE_FORMAT
            longest = max(longest, max((len(p) for p in display_text(cell, c).split("\n")), default=0))
        widths[c] = min(longest + 2, max_width)
        ws.column_dimensions[get_column_letter(c)].width = widths[c]

    for r in range(1, nrows + 1):
        lines = 1
        for c in range(1, ncols + 1):
            text = display_text(ws.cell(row=r, column=c), c)
            usable = max(1, widths[c] - 1)
            lines = max(lines, sum(max(1, math.ceil(len(p) / usable)) for p in text.split("\n")))
        ws.row_dimensions[r].height = lines * 15


def _write_report(out_path: Path, filtered: pd.DataFrame, sequences: pd.DataFrame) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = config.SHEET_FILTERED
    _write_filtered_sheet(ws, filtered)
    _write_sequences_sheet(wb.create_sheet(config.SHEET_CONTINUOUS), sequences)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def run_pipeline(raw_path, calendar_path, out_path, log=print) -> ReportSummary:
    """Build the two-sheet report. `log` receives progress messages (GUI or print)."""
    raw_path, calendar_path, out_path = Path(raw_path), Path(calendar_path), Path(out_path)

    log(f"Reading raw data:   {raw_path}")
    raw = pd.read_excel(raw_path)  # read once, shared by both sheets
    filtered = shaping.transform(raw)
    log(f"  shaping step kept {len(filtered)} target-leave rows")
    records = load_records(raw)
    log(f"  loaded {len(records)} leave records")

    log(f"Reading calendar:   {calendar_path}")
    cal = WorkCalendar.from_xlsx(calendar_path)

    if cal.year is not None:
        before = len(records)
        records = [r for r in records if r.start_date.year == cal.year]
        if before != len(records):
            log(f"  filtered to {len(records)} records in calendar year {cal.year} "
                f"(dropped {before - len(records)})")

    sequences = detect_all(records, cal)
    log(f"Detected {len(sequences)} continuous leave sequences")

    _write_report(out_path, filtered, to_dataframe(sequences))
    log(f"Wrote output:       {out_path}")
    return ReportSummary(len(filtered), len(sequences))


def main(argv: list[str]) -> int:
    raw_path = Path(argv[1]) if len(argv) > 1 else _find_raw_file()
    out_path = Path(argv[2]) if len(argv) > 2 else config.DATA_OUTPUT / report_filename(raw_path)
    run_pipeline(raw_path, config.CALENDAR_FILE, out_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
