"""Convert a human-readable calendar markdown into an xlsx.

The generated workbook has columns: Date, Day, Status, Holiday or event.
Only `Date` and `Status` are required by the pipeline.

Usage:
    python src/convert_calendar.py [source.md] [output.xlsx]
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import config

# Matches: | 2026-01-01 | Thursday | Public holiday | New Year's Day |
_ROW_RE = re.compile(
    r"\|\s*(\d{4}-\d{2}-\d{2})\s*\|\s*([A-Za-z]+)\s*\|\s*([^|]+?)\s*\|\s*([^|]+?)\s*\|"
)
COLUMNS = ["Date", "Day", "Status", "Holiday or event"]


def parse_markdown(md_path: str | Path) -> pd.DataFrame:
    text = Path(md_path).read_text(encoding="utf-8")
    rows = [
        (pd.Timestamp(d), day.strip(), status.strip(), event.strip())
        for d, day, status, event in _ROW_RE.findall(text)
    ]
    return pd.DataFrame(rows, columns=COLUMNS).sort_values("Date").reset_index(drop=True)


def convert(md_path: str | Path, xlsx_path: str | Path) -> int:
    df = parse_markdown(md_path)
    xlsx_path = Path(xlsx_path)
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    df.to_excel(xlsx_path, index=False, engine="openpyxl")

    wb = load_workbook(xlsx_path)
    ws = wb.active
    for c in range(1, ws.max_column + 1):
        longest = len(str(ws.cell(row=1, column=c).value or ""))
        ws.cell(row=1, column=c).font = Font(bold=True)
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=c)
            if c == 1 and cell.value is not None:
                cell.number_format = "m/d/yyyy"
                text = f"{cell.value.month}/{cell.value.day}/{cell.value.year}"
            else:
                text = str(cell.value) if cell.value is not None else ""
            longest = max(longest, len(text))
        ws.column_dimensions[get_column_letter(c)].width = longest + 2
    wb.save(xlsx_path)
    return len(df)


def main(argv: list[str]) -> int:
    md_path = Path(argv[1]) if len(argv) > 1 else config.CALENDAR_MD
    xlsx_path = Path(argv[2]) if len(argv) > 2 else config.CALENDAR_FILE
    count = convert(md_path, xlsx_path)
    print(f"Converted {count} calendar rows -> {xlsx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
